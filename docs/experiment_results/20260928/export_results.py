"""只读汇总已验收实验，生成七类实验的 Excel；不运行训练或修改原始产物。

运行：/home/yihangjie/miniconda3/envs/torch/bin/python docs/experiment_results/20260928/export_results.py
"""
from __future__ import annotations

import csv
import hashlib
import json
import math
import statistics as st
from collections import Counter, defaultdict
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.utils import get_column_letter
from sensitivity_charts import build_sensitivity_charts

ROOT = Path(__file__).resolve().parents[3]
OUT = Path(__file__).resolve().parent
P31 = ROOT / 'logs/exp031/exp031_full_20260729_174215'
P32 = ROOT / 'logs/exp032/exp032_full_20260917_2020/execution_revisions/adaptive_pgd10_v1'
P33 = ROOT / 'logs/exp033/exp033_full_20260923_v6'
P34 = ROOT / 'logs/exp034/exp034_full_20260928_v1'
P35 = ROOT / 'logs/exp035/exp035_full_20260928_v1'
SENSITIVITY_SEEDS = list(range(42, 52))
SEEDS = [42, 43, 44, 45, 46]
DATASETS = ['thubenchmark', 'seediv', 'bciciv2a']
MODELS = ['eegnet', 'deepconvnet', 'tsception', 'atcnet', 'conformer', 'tcnet']
DS = dict(thubenchmark='THUBenchmark', seediv='SEED-IV', bciciv2a='bciciv2a')
MODEL = dict(zip(MODELS, ['EEGNet', 'DeepConvNet', 'TSCeption', 'ATCNet', 'Conformer', 'TCNet']))
METHOD = {
    'clean': 'clean-only', 'madry': 'Madry', 'trades': 'TRADES', 'fbf': 'FBF',
    'ea_forward': 'EA-forward（ABAT-style）', 'rpcf_at': 'CAF（RPCF_AT）',
    'clean_magnet': 'MagNet-Reformer + clean-only', 'clean_dcae': 'DCAE + clean-only',
    'madry_magnet': 'MagNet-Reformer + Madry', 'madry_dcae': 'DCAE + Madry',
    'madry_tnp_r25': 'TRP+Madry · r25', 'madry_tnp_r30': 'TRP+Madry · r30',
    'rpcf_at_tnp_r25': 'TRP+CAF · r25', 'rpcf_at_tnp_r30': 'TRP+CAF · r30',
    'ptr': 'PTR（当前 TNP/TRP）', 'tr_dense': '普通 TR', 'tt_dense': '普通 TT',
    'tr_time': '时间张量化 TR', 'tt_time': '时间张量化 TT', 'tucker': 'Tucker', 'svd': '矩阵 SVD',
    'caf': 'CAF（RPCF_AT）', 'trp_caf': 'TRP+CAF', 'trp_clean': 'TRP+clean',
    'trp_madry': 'TRP+Madry', 'raw': 'clean-only',
    'magnet': 'MagNet-Reformer（EEG adaptation）', 'dcae': 'DCAE',
}
METHODS = ['clean', 'madry', 'trades', 'fbf', 'ea_forward', 'clean_magnet',
           'clean_dcae', 'madry_magnet', 'madry_dcae', 'rpcf_at', 'madry_tnp_r25', 'madry_tnp_r30',
           'rpcf_at_tnp_r25', 'rpcf_at_tnp_r30']
ATTACK = {'pgd': 'PGD-200 L∞ ε=0.03', 'fgsm': 'FGSM L∞ ε=0.03',
          'pgd_l2': 'PGD-200 L2 ε=1', 'autoattack': 'AutoAttack L∞ ε=0.03',
          'cw': 'CW L2（无半径约束）', 'pgd10': 'PGD-10 L∞ ε=0.03'}
EVAL = {'classifier_whitebox': '分类器白盒', 'nonadaptive_purification': '非自适应净化',
        'adaptive_exact_gradient': '自适应：真实梯度', 'adaptive_bpda': '自适应：BPDA', 'timing': '独占 GPU 计时'}
EXPS = ['1_主实验', '2_跨攻击跨范数', '3_Adaptive_attack', '4_张量网络效果效率',
        '5_参数敏感性', '6_消融实验', '7_可视化']
sources, records, gaps, validations = {}, [], [], []
# 保持既有来源ID稳定，避免只更新敏感性时改变其他实验页的引用。
previous_audit = json.loads((OUT/'validation.json').read_text()) if (OUT/'validation.json').exists() else {}
previous_source_ids = {Path(r['absolute_path']).resolve(): r['id'] for r in previous_audit.get('sources', [])}
next_source_number = max((int(v.removeprefix('SRC')) for v in previous_source_ids.values()), default=0) + 1


def source(path):
    global next_source_number
    path = Path(path).resolve()
    if path not in sources:
        assert path.is_file(), path
        if path in previous_source_ids:
            sources[path] = previous_source_ids[path]
        else:
            sources[path] = f'SRC{next_source_number:03d}'
            next_source_number += 1
    return sources[path]


def read_csv(path):
    source(path)
    with open(path, newline='', encoding='utf-8-sig') as f:
        return [dict(r, _path=str(path), _line=i) for i, r in enumerate(csv.DictReader(f), 2)]


def read_json(path):
    source(path)
    return json.loads(Path(path).read_text())


def scalar(value):
    if value in ('', None):
        return None
    if isinstance(value, (list, dict)):
        return json.dumps(value, ensure_ascii=False)
    if isinstance(value, (int, float)):
        return value
    try:
        n = float(value)
        return n if math.isfinite(n) else None
    except (ValueError, TypeError):
        return value


METRICS = {
    'standard_accuracy': ('SA 干净准确率', '%'), 'robust_accuracy': ('RA 鲁棒准确率', '%'),
    'decomposition_seconds_mean': ('分解耗时（样本均值）', 's'),
    'total_seconds_mean': ('总净化耗时（样本均值）', 's'),
    'total_seconds_median': ('总净化耗时（样本中位数）', 's'),
    'peak_memory_mib': ('峰值显存', 'MiB'), 'parameters': ('分解参数量', '个'),
    'native_compression_ratio': ('原始 EEG 压缩率', '倍'),
    'tensor_compression_ratio': ('张量表示压缩率', '倍'),
    'budget_deviation': ('参数预算偏差', '%'),
    'clean_mse_mean': ('净化 clean 对原始 clean 的 MSE', 'MSE'),
    'adv_mse_to_clean_mean': ('净化 adv 对原始 clean 的 MSE', 'MSE'),
    'removed_mse_mean': ('净化 adv 对净化前 adv 的 MSE', 'MSE'),
    'attack_mse': ('攻击扰动 MSE', 'MSE'), 'purified_clean_loss': ('净化 clean loss', 'loss'),
    'bpda_purified_adv_loss': ('BPDA 净化 adv loss', 'loss'),
}
for domain in ('clean', 'adv'):
    for mode in ('changed', 'corrected', 'damaged'):
        METRICS[f'{domain}_{mode}_fraction'] = (f'{domain} {mode} 比例', '%')


def record(exp, r, **extra):
    d = {k: v for k, v in r.items() if k not in ('clean_predictions', 'adv_predictions',
                                               'clean_confidence', 'adv_confidence')}
    d.update(extra)
    d['exp'] = exp
    d['seed'] = int(d['seed'])
    d['sample_num'] = int(d.get('sample_num') or 512)
    d['fold'] = int(d.get('fold') or 0)
    d['scope'] = d.get('scope') or 'S'
    d['source_id'] = source(d['_path'])
    d.setdefault('config', '')
    d.setdefault('evaluation', '')
    for key in METRICS:
        if key in d:
            d[key] = scalar(d[key])
    for k in ('standard_accuracy', 'robust_accuracy'):
        if d.get(k) is not None:
            assert 0 <= d[k] <= 1, (d, k)
    assert d['seed'] in (SENSITIVITY_SEEDS if exp == EXPS[4] else SEEDS)
    records.append(d)
    return d


def missing(exp, ds, model, method, config, reason):
    gaps.append([exp, DS.get(ds, ds), MODEL.get(model, model), METHOD.get(method, method),
                 config, 'Pending / 缺少结果', reason])


# 正式验收与来源冻结。
c32 = read_json(P32 / 'summary/completeness.json')
c33 = read_json(P33 / 'summary/report.json')
assert c32['completed'] and not c32['errors'] and c32['metric_rows'] == 5580
assert c33['complete'] and not c33['errors'] and c33['completed_task_count'] == 395
r32 = read_csv(P32 / 'summary/conditions_long.csv')
r33 = read_csv(P33 / 'summary/metrics_long.csv')
assert len(r32) == 5580 and len(r33) == 465
for r in r32:
    assert r['sample_num'] == '512' and r['scope'] == 'S'
    config = 'r' + r['method'].split('_r')[-1] if '_tnp_r' in r['method'] else '无测试净化 rank'
    if r['attack'] == 'pgd':
        record(EXPS[0], r, config=config)
    if r['dataset'] == 'thubenchmark' and r['model'] == 'eegnet':
        if r['attack'] in ('fgsm', 'pgd', 'pgd_l2', 'autoattack', 'cw'):
            record(EXPS[1], r, config=config)
        elif r['attack'] == 'pgd10' and r['evaluation'] == 'adaptive_exact_gradient':
            record(EXPS[2], r, config=config)
assert len([r for r in records if r['exp'] == EXPS[0]]) == 1080
assert len([r for r in records if r['exp'] == EXPS[1]]) == 300

bpda_path = P31 / 'summary_thu_eegnet/bpda.csv'
bpda = read_csv(bpda_path)
assert len(bpda) == 10
for r in bpda:
    record(EXPS[2], r, dataset='thubenchmark', model='eegnet',
           method='rpcf_at_tnp_r' + r['rank'], attack='pgd10', evaluation='adaptive_bpda',
           standard_accuracy=r['purified_clean_accuracy'], robust_accuracy=r['bpda_purified_adv_accuracy'],
           clean_mse_mean=r['mean_clean_mse'], adv_mse_to_clean_mean=r['mean_adv_mse'], config='r' + r['rank'])

DEFAULTS = dict(clean_ce_weight=1., pur_ce_weight=.5, adv_pur_ce_weight=1., lambda_pur=.2, lambda_adv_pur=.5)
# EXP-035的正式合并表替换原275条敏感性，不能追加后重复计入42–46。
c35 = read_json(P35/'summary/report.json')
assert c35['complete'] and not c35['smoke'] and not c35['errors'] and not c35['pending_tasks']
assert c35['completed_task_count'] == c35['task_count'] == 345
assert c35['metric_rows'] == 550 and c35['new_metric_rows'] == 275 and c35['condition_count'] == 55
r35 = read_csv(P35/'summary/metrics_long.csv')
assert len(r35) == 550 and Counter(int(r['seed']) for r in r35) == {seed:55 for seed in SENSITIVITY_SEEDS}
old_sensitive = {(r['task_id'], r['method'], r['rank']): r for r in r33 if r['group'] in ('rank','loss')}
for r in r35:
    assert r['group'] in ('rank','loss') and r['sample_num'] == '512'
    if int(r['seed']) in SEEDS:
        old = old_sensitive[(r['task_id'],r['method'],r['rank'])]
        assert all(float(r[k]) == float(old[k]) for k in ('standard_accuracy','robust_accuracy'))
for r in [r for r in r33 if r['group'] not in ('rank','loss')] + r35:
    g = r['group']
    exp = {'structure': EXPS[3], 'rank': EXPS[4], 'loss': EXPS[4], 'ablation': EXPS[5], 'visualize': EXPS[6]}[g]
    rank = r.get('rank', '')
    config = f'rank={rank or "无"}'
    if g == 'structure':
        config = f'预算档={r["budget_rank"]}'
    elif g == 'loss':
        param = r['scan_parameter'] or 'default'
        config = f'CE/KL | {param}={r["scan_value"] or "默认值"} | rank={rank}'
    elif g == 'rank':
        config = 'rank 扫描 | ' + config
    d = record(exp, r, config=config)
    if g in ('rank','loss'):
        d['task_source_id'] = source(ROOT/r['source_path'])
    if g == 'structure':
        d['model'] = d.get('model') or 'eegnet'
        d['dataset'] = d.get('dataset') or 'thubenchmark'
        d['attack'] = d.get('attack') or 'pgd'
        if r['peak_memory_bytes']:
            d['peak_memory_mib'] = float(r['peak_memory_bytes']) / 1024**2
        if r['task_id'].startswith('structure_'):
            p = P33 / 'metrics' / (r['task_id'] + '.json')
            task = read_json(p)
            diag = task['sample_diagnostics']
            assert len(diag) == 512
            for src, dst in [('clean_mse', 'clean_mse_mean'), ('adv_mse_to_clean', 'adv_mse_to_clean_mean'), ('removed_mse', 'removed_mse_mean')]:
                d[dst] = st.mean(x[src] for x in diag)
            d['diagnostics_source'] = source(p)
        if r['actual_rank']:
            d['actual_rank'] = r['actual_rank']
    if g == 'visualize' and not d['evaluation']:
        d['evaluation'] = 'classifier_whitebox' if d['method'] == 'raw' else 'nonadaptive_purification'

# EXP-034补充结果按实验用途分发；主实验PGD行同时用于跨攻击页，不当作额外独立seed。
c34 = read_json(P34 / 'summary/report.json')
assert c34['complete'] and not c34['smoke'] and not c34['errors'] and not c34['pending_tasks']
assert c34['task_count'] == c34['completed_task_count'] == 140
assert c34['metric_rows'] == c34['expected_metric_rows'] == 275
manifest34 = read_json(P34 / 'manifest.json')
r34 = read_csv(P34 / 'summary/metrics_long.csv')
assert len(r34) == 275
section_map = dict(main=EXPS[0], cross_attack=EXPS[1], adaptive=EXPS[2], structure=EXPS[3])
for r in r34:
    extra = dict(config=f'预算档={r["budget_rank"]}' if r['method'] == 'tr_time' else '无测试净化 rank')
    extra['task_source_id'] = source(ROOT / r['source_path'])
    extra['input_coordinate_system'] = '标准化EEG'
    if r['task_id'].startswith('raw_'):
        task = read_json(ROOT / r['source_path'])
        extra['input_coordinate_system'] = task['input_coordinate_system']
    if r['method'] in ('madry_magnet', 'madry_dcae'):
        extra['clean_evaluation_batch_size'] = manifest34['evaluation_batch_size']
    if r['peak_memory_bytes']:
        extra['peak_memory_mib'] = float(r['peak_memory_bytes']) / 1024**2
    for section in json.loads(r['sections']):
        record(section_map[section], r, **extra)
assert len([r for r in records if r['exp'] == EXPS[0]]) == 1260
assert len([r for r in records if r['exp'] == EXPS[1]]) == 350

# 各实验按各自seed范围独立分组，决不跨攻击、预算或口径混合。
group_fields = ['exp', 'dataset', 'model', 'method', 'attack', 'evaluation', 'config']
groups = defaultdict(list)
for d in records:
    groups[tuple(d.get(k, '') for k in group_fields)].append(d)
for key, vals in groups.items():
    expected_seeds = SENSITIVITY_SEEDS if key[0] == EXPS[4] else SEEDS
    assert sorted(v['seed'] for v in vals) == expected_seeds, (key, [v['seed'] for v in vals])
    assert len(set((v['scope'], v['sample_num']) for v in vals)) == 1, key
validations.append(f'{len(groups)}个条件中，55个敏感性条件精确覆盖42–51，其余条件精确覆盖42–46；未跨scope或样本数聚合。')
o35_rows = read_csv(P35/'summary/metrics_grouped.csv')
keys35 = ['group','method','rank','variant','scan_parameter','scan_value','evaluation','dataset','model','attack']
o35 = {tuple(r.get(k,'') for k in keys35):r for r in o35_rows}
assert len(o35) == 55
checked35 = 0
for vals in groups.values():
    if vals[0]['exp'] != EXPS[4]: continue
    o = o35[tuple(vals[0].get(k,'') for k in keys35)]
    assert int(o['n']) == 10 and json.loads(o['seeds']) == SENSITIVITY_SEEDS and o['complete'] == 'True'
    for metric in ('standard_accuracy','robust_accuracy'):
        nums = [v[metric] for v in vals]
        assert math.isclose(st.mean(nums),float(o[metric+'_mean']),abs_tol=1e-12,rel_tol=0)
        assert math.isclose(st.stdev(nums),float(o[metric+'_std']),abs_tol=1e-12,rel_tol=0)
        checked35 += 1
assert checked35 == 110
validations.append('EXP-035 550条逐seed数据、55条件共110对SA/RA均值及样本SD与正式聚合一致；原42–46数值逐项保留。')

# 与仓库正式聚合逐项交叉核验（均值/标准差单位在此转换）。
official = read_csv(P32 / 'summary/five_seed_mean_std.csv')
o32 = {(r['dataset'], r['model'], r['method'], r['attack'], r['evaluation']): r for r in official}
for key, vals in groups.items():
    if Path(vals[0]['_path']) != P32 / 'summary/conditions_long.csv':
        continue
    o = o32[tuple(vals[0][k] for k in ['dataset', 'model', 'method', 'attack', 'evaluation'])]
    for metric in ('standard_accuracy', 'robust_accuracy'):
        nums = [v[metric] for v in vals]
        assert abs(st.mean(nums)*100-float(o[metric+'_mean_percent'])) < 1e-8
        assert abs(st.stdev(nums)*100-float(o[metric+'_std_percent'])) < 1e-8
validations.append('EXP-032 所有采用条件的均值和 ddof=1 标准差均与正式五种子表一致。')

# EXP-034每个采用数值均与正式聚合表校验；同一PGD在两页出现时分别验证但不增加实验计数。
official34 = read_csv(P34 / 'summary/metrics_grouped.csv')
keys34 = ['dataset', 'model', 'method', 'attack', 'evaluation', 'budget_rank']
o34 = {tuple(r[k] for k in keys34): r for r in official34}
assert len(o34) == 55
checked34 = 0
for vals in groups.values():
    if Path(vals[0]['_path']) != P34 / 'summary/metrics_long.csv':
        continue
    o = o34[tuple(vals[0][k] for k in keys34)]
    assert int(o['seed_count']) == 5 and json.loads(o['seeds']) == SEEDS
    for metric in METRICS:
        src_metric = 'peak_memory_bytes' if metric == 'peak_memory_mib' else metric
        if not o.get(src_metric + '_mean'):
            continue
        scale = 1024**2 if metric == 'peak_memory_mib' else 1
        nums = [r[metric] for r in vals]
        assert math.isclose(st.mean(nums)*scale, float(o[src_metric+'_mean']), rel_tol=1e-10, abs_tol=1e-10), (metric, o)
        assert math.isclose(st.stdev(nums)*scale, float(o[src_metric+'_sample_std']), rel_tol=1e-10, abs_tol=1e-10), (metric, o)
        checked34 += 1
validations.append(f'EXP-034共275条独立seed结果；纳入各页57个条件，{checked34}项指标均值/样本SD与正式聚合一致。')

wb = Workbook()
wb.remove(wb.active)
NAVY, TEAL, LIGHT = '17365D', '087E8B', 'E8F4F5'
ORANGE, RED = 'FFF0CC', 'FCE8E6'
table_index = 0


def sheet(name, title, notes, headers, rows, percent_cols=()):
    global table_index
    ws = wb.create_sheet(name)
    end = max(len(headers), 10)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=min(end, 14))
    ws.cell(1, 1, title).font = Font(name='Calibri', size=18, bold=True, color='FFFFFF')
    ws.cell(1, 1).fill = PatternFill('solid', fgColor=NAVY)
    ws.row_dimensions[1].height = 34
    for i, note in enumerate(notes[:3], 2):
        ws.merge_cells(start_row=i, start_column=1, end_row=i, end_column=min(end, 14))
        c = ws.cell(i, 1, note)
        c.font = Font(name='Calibri', size=11, color='444444')
        c.alignment = Alignment(wrap_text=True, vertical='center')
        ws.row_dimensions[i].height = 30 if len(note) > 110 else 23
    header_row = 6
    for c, h in enumerate(headers, 1):
        x = ws.cell(header_row, c, h)
        x.font = Font(name='Calibri', color='FFFFFF', bold=True)
        x.fill = PatternFill('solid', fgColor=TEAL)
        x.alignment = Alignment(wrap_text=True, vertical='center')
        width = 20
        if any(t in h for t in ('方法', '配置', '说明', '来源', '路径', '原因', '攻击', '指标', '注意')):
            width = 35
        if h in ('seed', 'n_seed', '样本数', 'fold', '来源ID', '字节数'):
            width = 10
        if h in ('说明', '原因', '定义与注意事项', '实际rank（逐seed）', 'SHA256'):
            width = 70
        if h in ('仓库相对路径', '绝对路径'):
            width = 85
        ws.column_dimensions[get_column_letter(c)].width = width
    ws.row_dimensions[header_row].height = 32
    for ri, row in enumerate(rows, 7):
        ws.row_dimensions[ri].height = 34
        for ci, value in enumerate(row, 1):
            if isinstance(value, float):
                assert math.isfinite(value), (name, ri, ci)
            cell = ws.cell(ri, ci, value)
            cell.font = Font(name='Calibri', size=11, color='233247')
            cell.alignment = Alignment(vertical='center', wrap_text=True)
            if ci in percent_cols and isinstance(value, (float, int)):
                cell.number_format = '0.00%'
            elif isinstance(value, float):
                cell.number_format = '0.000000' if abs(value) < 0.01 else '0.0000'
            if isinstance(value, str) and ('Pending' in value or '缺少' in value):
                cell.fill = PatternFill('solid', fgColor=RED)
    if rows:
        table_index += 1
        tab = Table(displayName=f'Results{table_index}', ref=f'A6:{get_column_letter(len(headers))}{6+len(rows)}')
        tab.tableStyleInfo = TableStyleInfo(name='TableStyleMedium2', showRowStripes=True)
        ws.add_table(tab)
    ws.freeze_panes = 'D7'
    ws.sheet_view.showGridLines = False
    ws.sheet_view.zoomScale = 85
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.orientation = 'landscape'
    ws.page_setup.paperSize = ws.PAPERSIZE_A3
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.print_title_rows = '1:6'
    ws.sheet_properties.tabColor = TEAL
    return ws


def stats(vals, metric):
    ns = [x[metric] for x in vals if isinstance(x.get(metric), (int, float))]
    return (st.mean(ns), st.stdev(ns) if len(ns) > 1 else None) if ns else (None, None)


def pm(vals, metric, scale=100):
    mean, std = stats(vals, metric)
    return f'{mean*scale:.2f} ± {std*scale:.2f}' if mean is not None and std is not None else None


def note_for(r):
    notes = []
    if r['method'] == 'ea_forward':
        notes.append('历史ABAT/EA-in-forward路径（非严格原论文复现）；EA有batch依赖/历史clean漂移')
        if r.get('attack') == 'pgd10' and r.get('task_source_id'):
            notes.append('ε在EA前原始缓存输入坐标施加；SA batch32，攻击/adv batch1，保留原subject-aware forward；与标准化输入攻击需区分')
    if r['method'] in ('clean_magnet', 'madry_magnet', 'magnet'):
        notes.append('仅 Reformer，EEG adaptation；不包含 detector/rejection')
    if r.get('evaluation') == 'nonadaptive_purification':
        notes.append('先攻击分类器再净化，非完整防御自适应鲁棒性')
    if r.get('evaluation') == 'adaptive_bpda':
        notes.append('EXP-031历史RPCF_AT+TNP；真实净化forward，identity backward；独立攻击并非PGD-200')
    if r.get('evaluation') == 'adaptive_exact_gradient':
        notes.append('可微净化器真实梯度，不称BPDA；PGD10其余参数与历史协议对齐')
    if r.get('attack') == 'pgd10' and r.get('evaluation') == 'classifier_whitebox':
        notes.append('原始分类器真实梯度PGD-10，无净化器；SA batch32，攻击/adv batch1')
    if r['method'] in ('madry_magnet', 'madry_dcae'):
        notes.append('固定Madry分类器；复用既有净化器权重，不对Madry重新训练；SA batch8')
    if r.get('attack') == 'cw':
        notes.append('无固定L2半径；不是ε=0.03约束的CW')
    return '；'.join(notes)


ACC_HEADERS = ['数据集', '模型', '方法', '配置 / rank', '攻击', '评估方式', 'n_seed', 'SA 均值±SD (%)',
               'RA 均值±SD (%)', 'SA mean', 'SA sample SD', 'RA mean', 'RA sample SD',
               '样本数/seed', '状态', '原始方法标签', '来源ID', '说明']


def acc_row(vals):
    r = vals[0]
    sa, sas = stats(vals, 'standard_accuracy')
    ra, ras = stats(vals, 'robust_accuracy')
    return [DS.get(r['dataset'], r['dataset']), MODEL.get(r['model'], r['model']), METHOD.get(r['method'], r['method']),
            r['config'], ATTACK.get(r['attack'], r['attack']), EVAL.get(r['evaluation'], r['evaluation']), len(vals),
            pm(vals, 'standard_accuracy'), pm(vals, 'robust_accuracy'), sa, sas, ra, ras,
            r['sample_num'], f'Complete · {len(vals)}/{len(vals)}', r['method'], ', '.join(sorted(set(v['source_id'] for v in vals))), note_for(r)]


def groups_for(exp):
    return [v for k, v in groups.items() if k[0] == exp]


def sortkey(v):
    r = v[0]
    return (DATASETS.index(r['dataset']), MODELS.index(r['model']),
            list(ATTACK).index(r['attack']), METHODS.index(r['method']) if r['method'] in METHODS else 99,
            r['method'], r['config'], r['evaluation'])


for exp, title, note in [
    (EXPS[0], '主实验 · 三数据集 × 六模型 × 全部已实现 baseline', 'SEED 已按用户确认采用 SEED-IV。PGD-200 L∞ ε=0.03；14个方法/配置，252组，1260条逐seed结果；已补AP+Madry。'),
    (EXPS[1], '跨攻击、跨范数 · THUBenchmark / EEGNet', 'FGSM / PGD-200 L∞ / PGD-200 L2(半径1) / AutoAttack / CW；14个方法/配置 × 5攻击 × 5seed，共70组。'),
]:
    sheet(exp, title, [note, '主表统一使用固定子集 S=512、fold0；五 seed=42–46，统计为均值±样本标准差(ddof=1)，不混用全测试集结果。',
          '包含 clean-only、Madry、TRADES、FBF、EA-forward(ABAT-style)、MagNet/DCAE各配clean-only和Madry、CAF，以及 Madry/CAF+TRP r25/r30；逐seed值见“逐seed明细”。'],
          ACC_HEADERS, [acc_row(v) for v in sorted(groups_for(exp), key=sortkey)], percent_cols=(10, 11, 12, 13))

adaptive_rows = []
adaptive_available = {v[0]['method']: v for v in groups_for(EXPS[2])}
for method in METHODS:
    if method in adaptive_available:
        adaptive_rows.append(acc_row(adaptive_available[method]))
    else:
        reason = '未找到该方法 PGD-10 完整防御攻击的五seed结果；不以非自适应 PGD-200 代替。'
        missing(EXPS[2], 'thubenchmark', 'eegnet', method, 'PGD-10 L∞ ε=0.03', reason)
        adaptive_rows.append(['THUBenchmark', 'EEGNet', METHOD[method], '固定r25/r30或无净化', ATTACK['pgd10'],
                              '待补充完整防御攻击', 0, None, None, None, None, None, None, 512,
                              'Pending / 缺少结果', method, None, reason])
ws = sheet(EXPS[2], 'Adaptive attack · AP 与 ABAT 路径重点标注',
           ['已有11/14组：5个原始分类器PGD-10、4个AP组合真实梯度、2个TRP+CAF BPDA；余CAF、TRP+Madry r25/r30待补。',
            '攻击协议：PGD-10 / L∞ ε=0.03 / α=0.006 / 攻击batch1 / clean起点 / 无额外重启与EOT / 最后一步；SA批次见备注。',
            '蓝色为AP组合，橙色为EA-forward(ABAT-style)；EA在EA前缓存输入坐标施加ε，其他方法在标准化输入施加ε，比较需注意。'],
           ACC_HEADERS, adaptive_rows, percent_cols=(10, 11, 12, 13))
for ri, row in enumerate(adaptive_rows, 7):
    if row[15] in ('clean_magnet', 'clean_dcae', 'madry_magnet', 'madry_dcae', 'ea_forward'):
        for c in ws[ri]:
            c.fill = PatternFill('solid', fgColor=ORANGE if row[15] == 'ea_forward' else LIGHT)

# 结构比较：效果和计时并列，每个计时 seed 仅16个样本，避免误称512。
struct_methods = ['ptr', 'tr_dense', 'tr_time', 'tt_dense', 'tt_time', 'tucker', 'svd']
struct_rows = []
for model in MODELS:
    for method in struct_methods:
        if method == 'tr_time' and model != 'eegnet':
            continue  # 用户明确将EXP-034时间TR限定为EEGNet，不新增范围外缺失项。
        for budget in ('25', '30'):
            vals = [r for r in records if r['exp'] == EXPS[3] and r['model'] == model and r['method'] == method and r['budget_rank'] == budget and r['evaluation'] != 'timing']
            times = [r for r in records if r['exp'] == EXPS[3] and r['model'] == model and r['method'] == method and r['budget_rank'] == budget and r['evaluation'] == 'timing']
            if not vals:
                reason = 'EXP-033结构比较只执行EEGNet；此模型未测TT/Tucker/SVD及同预算对照。'
                missing(EXPS[3], 'thubenchmark', model, method, f'预算档{budget}', reason)
                struct_rows.append(['THUBenchmark', MODEL[model], METHOD[method], int(budget), 0, None, None] + [None]*15 + ['Pending / 缺少结果', None, reason])
                continue
            assert len(vals) == len(times) == 5
            ranks = ' | '.join(f"{r['seed']}:{r.get('actual_rank') or 'PTR rank '+budget}" for r in vals)
            srcs = sorted(set(r['source_id'] for r in vals+times) | {r['diagnostics_source'] for r in vals if 'diagnostics_source' in r})
            struct_rows.append(['THUBenchmark', MODEL[model], METHOD[method], int(budget), 5,
                pm(vals, 'standard_accuracy'), pm(vals, 'robust_accuracy'),
                stats(times, 'decomposition_seconds_mean')[0], stats(times, 'decomposition_seconds_mean')[1],
                stats(times, 'total_seconds_mean')[0], stats(times, 'total_seconds_mean')[1],
                stats(times, 'total_seconds_median')[0], stats(times, 'peak_memory_mib')[0], stats(times, 'peak_memory_mib')[1],
                stats(times, 'parameters')[0], stats(times, 'native_compression_ratio')[0], stats(times, 'tensor_compression_ratio')[0],
                stats(vals, 'clean_mse_mean')[0], stats(vals, 'adv_mse_to_clean_mean')[0], stats(vals, 'removed_mse_mean')[0],
                stats(vals, 'budget_deviation')[0], ranks, 'Complete · 5/5', ', '.join(srcs),
                'Madry固定分类器/攻击；计时16样本(seed内8clean+8adv)，RTX3080独占；普通TR低预算允许例外' + ('；PTR复用结果未提供同格式MSE，留空' if method == 'ptr' else '')])
struct_headers = ['数据集', '模型', '分解方法', '预算档（非共同rank）', 'n_seed', 'SA 均值±SD (%)', 'RA 均值±SD (%)',
                  '分解时间 mean(s)', '分解时间 SD(s)', '总净化时间 mean(s)', '总净化时间 SD(s)', '各seed样本中位数的均值(s)',
                  '峰值显存 mean(MiB)', '峰值显存 SD(MiB)', '参数量 mean', '原始EEG压缩率 mean', '张量压缩率 mean',
                  'clean MSE mean', 'adv→clean MSE mean', 'removed MSE mean', '预算偏差 mean', '实际rank（逐seed）', '状态', '来源ID', '说明']
sheet(EXPS[3], '张量网络效果与效率 · 六模型覆盖检查',
      ['已有EEGNet：PTR / 普通TR / 时间TR / 普通TT / 时间TT / Tucker / SVD，各两预算、5seed；其余5模型的原六分解仍Pending。',
       '攻击PGD-200 L∞ ε=0.03，效果样本512/seed；预算25/30分别对应14,731/19,891参数，不代表所有分解使用相同rank。',
       '时间TR按用户要求仅THU/EEGNet；标准闭环TR，先validation选秩。耗时先seed内统计再跨seed均值/SD；效果512样本，计时16样本。'],
      struct_headers, struct_rows, percent_cols=(21,))

param_rows = []
for v in sorted(groups_for(EXPS[4]), key=lambda v: (v[0]['group'], v[0].get('scan_parameter',''), float(v[0].get('scan_value') or -1), v[0]['method'], float(v[0].get('rank') or 0))):
    r = v[0]
    param = r.get('scan_parameter') or ('rank' if r['group'] == 'rank' else 'default')
    val = scalar(r.get('scan_value'))
    default = DEFAULTS.get(param)
    factor = val/default if isinstance(val, (float, int)) and default else None
    param_rows.append(acc_row(v) + [param, val, default, factor, 'KL' if param.startswith('lambda_') else 'CE' if param.endswith('ce_weight') else 'rank/default'])
parameter_ws = sheet(EXPS[4], '参数敏感性 · rank 与 CE / KL 权重',
      ['rank=15,20,25,30,35,40；CAF原始分类器作为参考。权重逐项扫描0×/0.5×/2×，并保留默认值参考；净化rank25/30分开。',
       '默认值：clean_ce=1，pur_ce=0.5，adv_pur_ce=1，lambda_pur(KL)=0.2，lambda_adv_pur(KL)=0.5。实际扫描值和倍率单独列出。',
       'THUBenchmark / EEGNet / PGD-200 L∞ ε=0.03 / S512 / seeds42–51；EXP-035的55组均为10seed；六张敏感性图从第65行开始，逐点数值见“参数图表数据”。'],
      ACC_HEADERS + ['扫描参数', '实际系数', '默认系数', '相对默认倍率', '项类型'], param_rows, percent_cols=(10,11,12,13))
sensitivity_figure_count, chart_point_count = build_sensitivity_charts(records, OUT, DEFAULTS, source, sheet, parameter_ws, seeds=SENSITIVITY_SEEDS)
sheet(EXPS[5], '消融实验 · clean-only / Madry / CAF / TRP组合',
      ['clean-only、Madry、CAF及TRP+clean / TRP+Madry / TRP+CAF；净化组合分别展示rank25与rank30，共9组×5seed。',
       'THUBenchmark / EEGNet / PGD-200 L∞ ε=0.03 / S512；攻击针对各自分类器，属于方法级比较。',
       'TRP/TNP表示测试时净化；CAF对应训练后的RPCF_AT；本页明确区分原始分类器与非自适应净化。'],
      ACC_HEADERS, [acc_row(v) for v in sorted(groups_for(EXPS[5]), key=sortkey)], percent_cols=(10,11,12,13))

visual_ws = sheet(EXPS[6], '可视化 · TNP 与 AP 方法',
      ['同一clean-only分类器及其PGD-200对抗样本；raw、TRP rank25/30、MagNet-Reformer、DCAE。上表为完整S512汇总，案例图在下方。',
       '20张原始案例图全部嵌入：5seed × 成功/失败/clean损伤/方法分歧；每图含clean、扰动、adv、净化/移除/残余信号，波形/PSD/STFT。',
       '案例按类别规则选取，仅用于定性展示，不替代五seed总体统计。原始PNG/PDF路径及source_index在图前列明。'],
      ACC_HEADERS, [acc_row(v) for v in sorted(groups_for(EXPS[6]), key=sortkey)], percent_cols=(10,11,12,13))
category = dict(success='成功恢复', failure='净化失败', clean_damage='干净样本受损', disagreement='方法分歧')
img_row, figure_count = 15, 0
for seed in SEEDS:
    cpath = P33 / f'figures/seed{seed}/seed{seed}_cases.json'
    cs = read_json(cpath)
    for case in cs['cases']:
        png = ROOT / (case['figure_stem'] + '.png')
        pdf = ROOT / (case['figure_stem'] + '.pdf')
        source(png); source(pdf)
        visual_ws.merge_cells(start_row=img_row, start_column=1, end_row=img_row, end_column=14)
        cell = visual_ws.cell(img_row, 1, f"Seed {seed} · {category[case['category']]} · source_index={case['source_index']} · label={case['label']} · channel={case['channel']} · fs={case['sampling_rate']}Hz")
        cell.font = Font(size=14, bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor=NAVY)
        visual_ws.row_dimensions[img_row].height = 30
        for offset, path in [(1, png), (2, pdf)]:
            visual_ws.merge_cells(start_row=img_row+offset, start_column=1, end_row=img_row+offset, end_column=14)
            cell = visual_ws.cell(img_row+offset, 1, str(path.relative_to(ROOT)))
            cell.hyperlink = str(path)
            cell.font = Font(color='0563C1', underline='single', size=10)
        im = XLImage(str(png))
        scale = 1320 / im.width
        im.width, im.height = 1320, im.height * scale
        visual_ws.add_image(im, f'A{img_row+4}')
        img_row += 6 + math.ceil(im.height / 20)
        figure_count += 1
assert figure_count == 20
visual_ws.freeze_panes = 'D7'
visual_ws.print_area = f'A1:N{img_row}'

# 可筛选的逐seed指标长表，所有数值带单位及CSV行号/JSON路径。
detail_rows = []
for r in records:
    for key, (label, unit) in METRICS.items():
        val = r.get(key)
        if not isinstance(val, (int, float)):
            continue
        src = r.get('diagnostics_source') if key in ('clean_mse_mean', 'adv_mse_to_clean_mean', 'removed_mse_mean') and r.get('diagnostics_source') else r['source_id']
        detail_rows.append([r['exp'], DS.get(r['dataset'],r['dataset']), MODEL.get(r['model'],r['model']),
                           METHOD.get(r['method'],r['method']), r['method'], r['config'], ATTACK.get(r['attack'],r['attack']),
                           EVAL.get(r['evaluation'],r['evaluation']), r['seed'], r['fold'], r['scope'], r['sample_num'],
                           label, key, val*100 if unit == '%' else val, unit, src,
                           r.get('_line') if src == r['source_id'] else 'sample_diagnostics → seed内均值',
                           r.get('task_source_id'), r.get('task_id'), r.get('classifier_method'),
                           r.get('clean_evaluation_batch_size'), r.get('evaluation_batch_size'),
                           r.get('input_coordinate_system'), r.get('attack_protocol')])
sheet('逐seed明细', '逐seed原始指标 · 可按实验、方法、指标筛选',
      ['准确率及比例在本页为0–100数值，单位%；MSE/秒/MiB/参数量各用明确单位。未包含巨大预测数组，其原始结果路径保留。',
       '每条指标可由来源ID定位到CSV行号或JSON逐样本诊断；重构MSE是512样本均值，计时指标来自每seed16样本。'],
      ['实验', '数据集', '模型', '方法', '原始方法标签', '配置', '攻击', '评估方式', 'seed', 'fold', 'scope', '样本数',
       '指标', '原始字段', '值', '单位', '来源ID', 'CSV行号 / JSON位置', '任务JSON来源ID',
       '任务ID', '分类器方法', 'clean推理batch', 'adv评估batch', '输入坐标', '攻击协议JSON'], detail_rows)

# 原有全测试集结果作为单独附录保存，避免主表遗漏历史数据或误混口径。
f31 = read_csv(P31 / 'summary/conditions_long.csv')
full_groups = defaultdict(list)
for r in f31:
    if r['attack'] == 'pgd' and r['rank'] == 'raw' and r['metric'] in ('standard_accuracy','robust_accuracy'):
        full_groups[(r['dataset'], r['model'], r['method'], r['metric'])].append(r)
full_rows = []
for (ds, model, method, metric), vals in sorted(full_groups.items()):
    assert sorted(int(v['seed']) for v in vals) == SEEDS
    nums = [float(v['value']) for v in sorted(vals,key=lambda x:int(x['seed']))]
    full_rows.append([DS[ds],MODEL[model],METHOD[method],METRICS[metric][0], st.mean(nums),st.stdev(nums),*nums,
                     'F：完整测试集', source(P31/'summary/conditions_long.csv'), '仅历史参考；不能与主表S512直接相减'])
sheet('附录_全测试集主实验', 'EXP-031 历史全测试集主实验 · 独立口径',
      ['PGD-200 L∞ ε=0.03；三数据集×六模型×五个原始训练baseline；每项五seed。',
       '此页仅完整测试集 F；主实验页采用统一 S512 才能包含净化方法和新增baseline。F与S禁止合并统计。'],
      ['数据集','模型','方法','指标','mean','sample SD','seed42','seed43','seed44','seed45','seed46','scope','来源ID','说明'],
      full_rows, percent_cols=tuple(range(5,12)))

# 保留参数/预算/协议缺口；原始clean-only+TNP暂缓不冒充已有全矩阵。
missing('补充范围说明', '全部三数据集', '全部六模型', 'trp_clean', 'EXP-032全部五攻击/全矩阵',
        'EXP-032原450个clean-only+TNP任务Deferred；EXP-033仅补THU/EEGNet/PGD，已单独收录消融和可视化。')
sheet('缺失与边界', '缺失结果与比较边界',
      ['Pending表示未找到满足所列配置的正式结果，数值单元格留空；不以0替代，也未启动新实验。',
       'ABAT显示为仓库现有EA-forward(ABAT-style)路径，AP类包含MagNet-Reformer与DCAE；严格原论文复现范围见来源文档。'],
      ['实验','数据集','模型','方法','配置','状态','原因'], gaps)

overview = [
    [EXPS[0], '3数据集×6模型×14方法/配置', 'PGD-200 L∞ .03', '252/252组，5/5 seeds', 'Complete', '统一S512；SEED-IV已获用户确认'],
    [EXPS[1], 'THU/EEGNet×14方法×5攻击', 'FGSM/PGD∞/PGD-L2/AA/CW', '70/70组，5/5 seeds', 'Complete', '跨范数分别展示，不计算混合均值'],
    [EXPS[2], 'THU/EEGNet；全部方法列出', 'PGD-10 L∞ .03', '11/14组有5seed结果', 'Partial / 缺少3组', '已补5原始分类器与AP+Madry；CAF和TRP+Madry仍Pending'],
    [EXPS[3], '原6模型×6分解；新增时间TR仅EEGNet', 'PGD-200 L∞ .03', 'EEGNet 14组效果+14组计时', 'Partial / 其余5模型Pending', '计时16样本/seed；效果512样本/seed'],
    [EXPS[4], 'THU/EEGNet；rank与CE/KL', 'PGD-200 L∞ .03', '55/55组，10/10 seeds', 'Complete', '6张SA/RA敏感性图+134个绘图点源数据；1×复用默认实验'],
    [EXPS[5], 'clean/Madry/CAF及TRP组合', 'PGD-200 L∞ .03', '9/9组，5/5 seeds', 'Complete', '三种净化组合各r25/r30'],
    [EXPS[6], 'TNP25/30 vs MagNet/DCAE', 'PGD-200 L∞ .03', '5组×5seed；20张嵌入图', 'Complete', '每seed四类案例；总体指标与案例分开'],
    ['逐seed明细', f'{len(detail_rows)}条指标', '各实验独立', '来源ID+CSV行号/JSON字段', '可追溯', '准确率/比例0–100%；其余明确单位'],
    ['附录_全测试集主实验', 'EXP-031原始F范围', 'PGD-200 L∞ .03', '180指标组×5seed', '参考', '不与S512净化结果混合'],
]
ws = sheet('说明与覆盖', 'EEG 鲁棒性实验结果汇总 · 更新至2026-09-30',
           ['只读整理已完成EXP-031/032/033/034/035；覆盖用户指定七类实验。每项均保留来源，未新增训练或攻击。',
            '统计：参数敏感性10seed=42–51；其余实验5seed=42–46。算术均值±样本SD(ddof=1)，fold0；SA/RA为干净/鲁棒准确率。',
            '主表百分比均值±SD为便于阅读的文本；旁边mean/SD为数值，逐seed值可在明细页筛选。点击实验名跳转。'],
           ['实验页','覆盖范围','攻击','完成情况','状态','说明'], overview)
for ri, row in enumerate(overview,7):
    if row[0] in wb.sheetnames:
        ws.cell(ri,1).hyperlink = f"#'{row[0]}'!A1"
        ws.cell(ri,1).font = Font(color='0563C1',underline='single')

protocols = [
    ['PGD-200 L∞', 'eps=.03, alpha=2/255, 200 steps, no random start', 'docs/EXP032_PLAN.md'],
    ['PGD-200 L2', '全样本绝对L2半径1.0，alpha=.1，200 steps，5 random restarts', 'docs/EXP032_PLAN.md'],
    ['FGSM', 'L∞ eps=.03，1 step', 'docs/EXP032_PLAN.md'],
    ['AutoAttack', 'standard L∞ eps=.03', 'docs/EXP032_PLAN.md'],
    ['CW', '无半径约束L2；steps200, lr=.1, c=10000, kappa=1', 'docs/EXP032_PLAN.md'],
    ['PGD-10 adaptive', 'L∞ eps=.03, alpha=.006, steps10, attack batch1, clean起点，无额外restart/EOT；最后一步。TNP BPDA、AP组合/原始分类器真实梯度', 'docs/EXP034_PLAN.md'],
    ['ABAT标签', '仓库历史ABAT EA-in-forward路径；表中原始method=ea_forward，显示ABAT-style，不宣称严格原论文复现', 'docs/EXPERIMENTS.md'],
    ['AP标签', '本整理指MagNet-Reformer(EEG adaptation)/DCAE净化方法；MagNet不含detector/rejection', 'docs/EXP032_BASELINES.md'],
    ['CAF / TRP', 'CAF=RPCF_AT训练分类器；TRP/TNP为测试净化；PTR为结构对照中的当前净化分解', 'docs/EXP033_PLAN.md'],
    ['子集S', 'stable_subset_indices；NumPy RandomState(seed+fold*1000)，每seed512；统一索引/标签', 'docs/EXP032_PLAN.md'],
    ['结构预算', '预算25/30目标14731/19891 params；普通TR低档13432为显式例外，actual_rank各方法定义不同', 'docs/EXP033_PLAN.md'],
    ['计时范围', '8clean+8adv，独占GPU、预热不计时；显存为torch max_memory_allocated，非整卡显存', 'rpcf/exp033_worker.py'],
    ['EXP-034 AP+Madry', 'MagNet/DCAE固定净化器配Madry；clean推理batch8，PGD10攻击/adv batch1。保留原clean-only搭配并分别列出', 'docs/EXP034_PLAN.md'],
    ['EXP-034 raw PGD10', '原始分类器clean推理batch32、攻击/adv batch1；EA-forward在EA前原始缓存输入坐标攻击并经subject-aware forward，其他方法标准化EEG', 'rpcf/exp034_worker.py'],
    ['时间张量化TR', '13阶标准闭环TR；validation32样本选秩；预算25/30实参14733/20325，均在±5%；仅THU/EEGNet，效果S512、计时16样本', 'docs/EXP034_PLAN.md'],
    ['参数敏感性图', 'EXP-035：55条件×10seed(42–51)；rank六点+CAF参考，五权重0/.5/1/2倍，1倍共用默认；均值±样本SD+逐seed散点', 'docs/EXP035_PLAN.md'],
    ['EA审计', 'EXP-031历史clean漂移65/90条件；EXP-032重新推理S子集并保留state/batch审计，数值未手动纠正', 'docs/EXP031_results_20260917_final/report.md'],
]
for _,_,path in protocols:
    source(ROOT/path)
source(P32/'summary/baseline_audit.csv')
source(P33/'summary/metrics_grouped.csv')
source(P31/'summary/completeness.json')
source(P31/'summary_thu_eegnet/bpda_summary.csv')
source(P33/'summary/paired_grouped.csv')
source(Path(__file__))
source(OUT/'sensitivity_charts.py')
sheet('协议与指标', '攻击、统计与方法命名口径',
      ['源代码标签均保留，展示名仅为便于阅读；不能将不同范数、梯度方式、样本范围或模型来源混为一个实验。'],
      ['项目','定义与注意事项','来源文件'], protocols)

source_rows = []
for p, sid in sources.items():
    raw = p.read_bytes()
    source_rows.append([sid, str(p.relative_to(ROOT)), str(p), len(raw),
                        datetime.fromtimestamp(p.stat().st_mtime).isoformat(timespec='seconds'), hashlib.sha256(raw).hexdigest()])
ws = sheet('来源索引', '来源索引 · 文件路径与 SHA256',
           ['实验CSV/JSON及原始案例图均只读；本次生成的敏感性PNG/PDF/CSV另行登记。CSV行号含表头，一基编号。',
            'Excel内图片自包含；源文件超链接需在原工作区访问。源SHA256可用于确认后续是否发生变化。'],
           ['来源ID','仓库相对路径','绝对路径','字节数','修改时间（本地）','SHA256'], source_rows)
for ri, row in enumerate(source_rows,7):
    ws.cell(ri,3).hyperlink = row[2]
    ws.cell(ri,3).font = Font(color='0563C1',underline='single')

# 七类实验排在总览之后，其余为审计附录。
order = ['说明与覆盖']+EXPS+['参数图表数据','逐seed明细','附录_全测试集主实验','缺失与边界','协议与指标','来源索引']
wb._sheets = [wb[n] for n in order]
wb.active = 0
wb.properties.title = 'EEG鲁棒性七类实验结果汇总'
wb.properties.subject = 'EXP-031至EXP-035；敏感性十seed、其余五seed可追溯结果'
wb.properties.creator = 'EEG_TNP research project'
OUT.mkdir(parents=True,exist_ok=True)
target = OUT/'EEG_完整实验结果_20260928.xlsx'
staging = target.with_name(target.stem + '.staging.xlsx')
wb.save(staging)

# 验证生成文件可重新打开，图片数量、非空数值及sheet结构完整。
check = load_workbook(staging, read_only=False, data_only=True)
assert check.sheetnames == order
assert len(check[EXPS[6]]._images) == 20
assert check[EXPS[0]].max_row == 258
assert check[EXPS[1]].max_row == 76
assert check[EXPS[2]].max_row == 20
assert len(check[EXPS[4]]._images) == 6
assert check["参数图表数据"].max_row == 140
assert len(gaps) == 64 and len(records) == 2425
assert len(groups_for(EXPS[4])) == 55 and len(groups_for(EXPS[5])) == 9
validations.append('Excel可重新打开；14个sheet、26张内嵌PNG（20案例+6敏感性）；主实验252组、跨攻击70组、参数55组、消融9组。')
validations.append('所有缺失条件保留Pending和空数值；原实验文件未写入。')
# 更新前后其他实验页的数值、格式、超链接和案例图片必须保持一致。
preserved_sheets = [EXPS[i] for i in (0,1,2,3,5,6)] + ['附录_全测试集主实验','缺失与边界']
if target.exists():
    old_book = load_workbook(target, data_only=True)
    def cells(ws):
        return {(c.coordinate):(c.value,c.number_format,c.hyperlink.target if c.hyperlink else None)
                for row in ws for c in row if c.value is not None}
    for name in preserved_sheets:
        assert cells(old_book[name]) == cells(check[name]), f'Unrelated sheet changed: {name}'
        assert set(map(str,old_book[name].merged_cells.ranges)) == set(map(str,check[name].merged_cells.ranges))
        assert [hashlib.sha256(im._data()).hexdigest() for im in old_book[name]._images] == [hashlib.sha256(im._data()).hexdigest() for im in check[name]._images]
    old_book.close()
assert len(detail_rows) == 5770
assert all(check[EXPS[4]].cell(row,7).value == 10 for row in range(7,62))
validations.append('8个非敏感性实验/附录sheet的单元格值、数值格式、超链接、合并区和20张案例图片与更新前一致。')
check.close()
staging.replace(target)
audit = {'created_at': datetime.now().isoformat(timespec='seconds'), 'file': str(target),
         'seed_range': SEEDS, 'sensitivity_seed_range': SENSITIVITY_SEEDS, 'datasets': DS, 'record_count': len(records), 'metric_detail_count': len(detail_rows),
         'group_counts': dict(Counter(k[0] for k in groups)), 'embedded_figures': figure_count + sensitivity_figure_count,
         'sensitivity_figures': sensitivity_figure_count, 'chart_data_points': chart_point_count,
         'exp035_combined_seed_rows': len(r35), 'exp035_adopted_statistic_pairs': checked35,
         'exp034_unique_seed_rows': len(r34), 'exp034_adopted_statistic_pairs': checked34,
         'missing_rows': len(gaps), 'validations': validations, 'source_count': len(sources),
         'sources': [dict(zip(['id','relative_path','absolute_path','bytes','mtime','sha256'],r)) for r in source_rows]}
(OUT/'validation.json').write_text(json.dumps(audit,ensure_ascii=False,indent=2)+'\n')
print(json.dumps({k:v for k,v in audit.items() if k != 'sources'},ensure_ascii=False,indent=2))
