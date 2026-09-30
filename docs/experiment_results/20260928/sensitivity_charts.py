"""从指定种子原始记录绘制敏感性图；默认点复用既有默认实验，不增加观测。"""
from __future__ import annotations

import csv
import math
import statistics as st
from pathlib import Path

import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Font, PatternFill


SEEDS = [42, 43, 44, 45, 46]
COLORS = ['#6B7280', '#2479B5', '#D67923']
VARIANTS = [('caf', 0, 'CAF raw'), ('trp_caf', 25, 'TRP+CAF r25'), ('trp_caf', 30, 'TRP+CAF r30')]
METRICS = [('standard_accuracy', 'Clean accuracy (SA, %)'), ('robust_accuracy', 'Robust accuracy (RA, %)')]


def build_sensitivity_charts(records, out, defaults, source, sheet, parameter_ws, seeds=SEEDS):
    """生成六张双面板图、逐点源数据，并把图与数据内嵌到工作簿。"""
    n = len(seeds)
    seed_label = f'{seeds[0]}-{seeds[-1]}'
    directory = out / 'sensitivity_figures'
    directory.mkdir(exist_ok=True)
    plt.rcParams.update({'font.family': 'DejaVu Sans', 'font.size': 10,
                         'axes.spines.top': False, 'axes.spines.right': False,
                         'pdf.fonttype': 42, 'savefig.facecolor': 'white'})
    raw = [r for r in records if r['exp'] == '5_参数敏感性']
    rows, figures = [], []

    def select(group, method, rank, param=None, factor=None):
        selected = [r for r in raw if r['group'] == group and r['method'] == method
                    and int(r.get('rank') or 0) == rank]
        if group == 'loss':
            if factor == 1:
                selected = [r for r in selected if not r.get('scan_parameter')]
            else:
                selected = [r for r in selected if r.get('scan_parameter') == param
                            and math.isclose(float(r['scan_value']), defaults[param] * factor)]
        selected.sort(key=lambda r: r['seed'])
        assert [r['seed'] for r in selected] == seeds, (group, method, rank, param, factor)
        return selected

    def point(figure, param, method, rank, factor, actual, metric, vals, reference=False):
        values = [r[metric] * 100 for r in vals]
        mean, sd = st.mean(values), st.stdev(values)
        rows.append([figure, param, method, rank, factor, actual, metric, mean, sd,
                     *values, '默认实验复用（1×）' if reference else '实测',
                     ', '.join(sorted({r['source_id'] for r in vals})),
                     '; '.join(f"seed{r['seed']}:{r['_line']}" for r in vals)])
        return mean, sd, values

    def decorate(ax, ylabel):
        ax.set_ylabel(ylabel)
        ax.grid(axis='y', color='#DCE3E9', linewidth=.7)
        ax.set_axisbelow(True)
        ax.tick_params(axis='both', labelsize=9)

    def save(fig, stem, caption):
        png, pdf = directory / (stem + '.png'), directory / (stem + '.pdf')
        fig.savefig(png, dpi=180)
        fig.savefig(pdf)
        plt.close(fig)
        source(png); source(pdf)
        figures.append((png, pdf, caption))

    ranks = [15, 20, 25, 30, 35, 40]
    fig, axes = plt.subplots(1, 2, figsize=(12, 4.5))
    fig.subplots_adjust(left=.07, right=.985, top=.78, bottom=.20, wspace=.24)
    fig.suptitle('Rank sensitivity', x=.07, y=.97, ha='left', fontsize=17, fontweight='bold')
    fig.text(.07, .88, f'THUBenchmark / EEGNet / PGD-200 L-infinity 0.03 / S512 / seeds {seed_label}', color='#475569')
    for ax, (metric, ylabel) in zip(axes, METRICS):
        vals = [point('rank', 'rank', 'trp_caf', rank, None, rank, metric,
                      select('rank', 'trp_caf', rank)) for rank in ranks]
        ax.errorbar(ranks, [x[0] for x in vals], yerr=[x[1] for x in vals],
                    color=COLORS[1], marker='o', capsize=4, linewidth=1.8, label='TRP+CAF')
        for x, (_, _, seedvals) in zip(ranks, vals):
            ax.scatter([x + (i-(n-1)/2)*.10 for i in range(n)], seedvals, color=COLORS[1], s=12, alpha=.38)
        mu, sd, _ = point('rank', 'rank', 'caf', 0, None, None, metric, select('rank', 'caf', 0))
        ax.axhline(mu, color=COLORS[0], linestyle='--', linewidth=1.3, label='CAF raw reference')
        ax.axhspan(mu-sd, mu+sd, color=COLORS[0], alpha=.10)
        ax.set_xticks(ranks)
        ax.set_xlabel('TRP rank')
        decorate(ax, ylabel)
    axes[0].legend(frameon=False, loc='best', fontsize=9)
    fig.text(.07, .05, f'Mean +/- sample SD (n={n}); dots show individual seeds. CAF reference band: mean +/- SD. Y-axes use data ranges.',
             fontsize=8.5, color='#475569')
    save(fig, 'rank_sensitivity', 'rank：TRP+CAF 15/20/25/30/35/40；灰色虚线和阴影为 CAF raw 均值±样本SD。')

    factors = [0., .5, 1., 2.]
    for param, default in defaults.items():
        fig, axes = plt.subplots(1, 2, figsize=(12, 4.7))
        fig.subplots_adjust(left=.07, right=.985, top=.73, bottom=.23, wspace=.24)
        kind = 'KL' if param.startswith('lambda_') else 'CE'
        fig.suptitle(f'{kind} weight sensitivity: {param}', x=.07, y=.97, ha='left', fontsize=16, fontweight='bold')
        fig.text(.07, .885, f'Default coefficient = {default:g} | THUBenchmark / EEGNet / PGD-200 L-infinity 0.03 / S512 / n={n}', color='#475569')
        for ax, (metric, ylabel) in zip(axes, METRICS):
            for vi, (method, rank, label) in enumerate(VARIANTS):
                vals = [point(param, param, method, rank, f, default*f, metric,
                              select('loss', method, rank, param, f), reference=(f == 1)) for f in factors]
                ax.errorbar(factors, [x[0] for x in vals], yerr=[x[1] for x in vals],
                            color=COLORS[vi], marker=['s', 'o', '^'][vi], capsize=3,
                            linewidth=1.5, label=label)
                for x, (_, _, seedvals) in zip(factors, vals):
                    ax.scatter([x+(i-(n-1)/2)*.004+(vi-1)*.017 for i in range(n)], seedvals,
                               color=COLORS[vi], s=10, alpha=.33)
            ax.axvline(1., color='#94A3B8', linestyle=':', linewidth=1)
            ax.set_xticks(factors, [f'{f:g}x\n({f*default:g})' for f in factors])
            ax.set_xlabel('Multiplier of default (actual coefficient)')
            decorate(ax, ylabel)
        handles, labels = axes[0].get_legend_handles_labels()
        fig.legend(handles, labels, loc='upper left', bbox_to_anchor=(.065, .835), ncol=3, frameon=False, fontsize=9)
        fig.text(.07, .045, f'Mean +/- sample SD; dots: seeds {seed_label}. The 1x points reuse each method\'s common default run. Y-axes use data ranges.',
                 fontsize=8.5, color='#475569')
        save(fig, param, f'{param}（{kind}）：默认值 {default:g}，横轴0/0.5/1/2倍；1倍复用各方法的默认实验。')

    headers = ['图名称', '扫描参数', '原始方法标签', 'rank', '相对默认倍率', '实际系数 / rank',
               '指标字段', 'mean (%)', 'sample SD (%)', *[f'seed{s} (%)' for s in seeds], '点来源类型', '来源ID', '逐seed CSV行号']
    data_path = directory / 'sensitivity_chart_data.csv'
    with data_path.open('w', encoding='utf-8-sig', newline='') as f:
        writer = csv.writer(f)
        writer.writerow(headers)
        writer.writerows(rows)
    source(data_path)
    sheet('参数图表数据', f'敏感性图表源数据 · 每个绘图点及{n}seed值',
          [f'图中均值与误差条为{n}seed算术均值±样本标准差(ddof=1)；数据单位全部为百分点（0–100）。',
           '权重1×从各方法/净化rank的共同默认实验复用，CSV行号重复属正常，不是额外独立重复实验；CAF raw为参考。',
           '六张图嵌在“5_参数敏感性”表格下方；同目录另附PNG/PDF和本页CSV。rank图的CAF参考不算rank=0扫描点。'],
          headers, rows)
    image_row = 65
    for png, pdf, caption in figures:
        parameter_ws.merge_cells(start_row=image_row, start_column=1, end_row=image_row, end_column=14)
        cell = parameter_ws.cell(image_row, 1, caption)
        cell.font = Font(size=13, bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor='17365D')
        parameter_ws.row_dimensions[image_row].height = 28
        for offset, path in [(1, png), (2, pdf)]:
            parameter_ws.merge_cells(start_row=image_row+offset, start_column=1, end_row=image_row+offset, end_column=14)
            c = parameter_ws.cell(image_row+offset, 1, f'{path.name}（点击打开原图）')
            c.hyperlink = str(path)
            c.font = Font(color='0563C1', underline='single')
        im = XLImage(str(png))
        im.height *= 1250 / im.width
        im.width = 1250
        parameter_ws.add_image(im, f'A{image_row+4}')
        image_row += 6 + math.ceil(im.height / 20)
    parameter_ws.print_area = f'A1:W{image_row}'
    assert len(figures) == 6 and len(rows) == 134
    return len(figures), len(rows)
